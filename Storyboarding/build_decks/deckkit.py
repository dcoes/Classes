"""Shared helpers for building CA 140 decks on the existing theme.

Method: copy a finished deck (theme carrier), strip its slides, add new ones.
This preserves theme, fonts, colour palette and layouts exactly, which is what
PROJECT.md requires -- decks are never regenerated from scratch.
"""
import copy
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

THEME_CARRIER = "Storyboarding/NewLectures/12_Animatics_REVISED.pptx"

# Layout indices (identical across 05 and 12)
L_TITLE = 0
L_TITLE_CONTENT = 1
L_SECTION = 2
L_TWO_CONTENT = 3
L_TITLE_ONLY = 5
L_BLANK = 6
L_TITLE_CAPTION = 10
L_QUOTE = 11

SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)


def new_deck(dest):
    """Copy the theme carrier to dest and remove all its slides."""
    shutil.copyfile(THEME_CARRIER, dest)
    prs = Presentation(dest)
    xml_slides = prs.slides._sldIdLst
    for sldId in list(xml_slides):
        rId = sldId.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        prs.part.drop_rel(rId)
        xml_slides.remove(sldId)
    return prs


def _strip_empty_placeholders(slide):
    """Remove placeholders left empty so they don't show prompt text."""
    for shape in list(slide.placeholders):
        if shape.placeholder_format.idx in (10, 11, 12):
            continue
        if not shape.has_text_frame:
            continue
        if not shape.text_frame.text.strip():
            shape._element.getparent().remove(shape._element)


def notes(slide, text):
    """Set speaker notes. Text is the verbatim script (CLAUDE.md 2)."""
    slide.notes_slide.notes_text_frame.text = text.strip()


def title_slide(prs, title, subtitle, note):
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    s.shapes.title.text = title
    s.placeholders[1].text = subtitle
    notes(s, note)
    return s


def section(prs, title, sub, note):
    s = prs.slides.add_slide(prs.slide_layouts[L_SECTION])
    s.shapes.title.text = title
    s.placeholders[1].text = sub
    notes(s, note)
    return s


def bullets(prs, title, items, note, layout=L_TITLE_CONTENT):
    """items: list of str, or (str, level) tuples."""
    s = prs.slides.add_slide(prs.slide_layouts[layout])
    s.shapes.title.text = title
    body = s.placeholders[1].text_frame
    body.clear()
    for i, item in enumerate(items):
        text, level = (item, 0) if isinstance(item, str) else item
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = text
        p.level = level
    notes(s, note)
    return s


def two_content(prs, title, left, right, note):
    s = prs.slides.add_slide(prs.slide_layouts[L_TWO_CONTENT])
    s.shapes.title.text = title
    for ph_idx, items in ((1, left), (2, right)):
        tf = s.placeholders[ph_idx].text_frame
        tf.clear()
        for i, item in enumerate(items):
            text, level = (item, 0) if isinstance(item, str) else item
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = text
            p.level = level
    notes(s, note)
    return s


def quote(prs, title, big, caption, note):
    s = prs.slides.add_slide(prs.slide_layouts[L_QUOTE])
    s.shapes.title.text = title
    s.placeholders[13].text = big
    s.placeholders[1].text = caption
    notes(s, note)
    return s


def title_only(prs, title, note):
    """Title-only slide, ready for shapes to be added underneath."""
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE_ONLY])
    s.shapes.title.text = title
    notes(s, note)
    return s


# ---------------------------------------------------------------- diagrams

INK = RGBColor(0x2B, 0x2B, 0x2B)
MUTE = RGBColor(0x8A, 0x8A, 0x8A)
ACCENT = RGBColor(0xC0, 0x50, 0x4D)
PAPER = RGBColor(0xF2, 0xF0, 0xEA)


def box(slide, x, y, w, h, text="", fill=None, line=INK, size=13,
        bold=False, shape=MSO_SHAPE.RECTANGLE, color=INK):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill if fill else PAPER
    sh.line.color.rgb = line
    sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return sh


def label(slide, x, y, w, text, size=13, bold=False, align=PP_ALIGN.CENTER,
          color=INK, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = align
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return tb


def line(slide, x1, y1, x2, y2, color=INK, width=1.25, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c


def arrow(slide, x, y, w, h, color=ACCENT, rotation=0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    if rotation:
        sh.rotation = rotation
    return sh


def figure(slide, cx, top, height, lean=0, color=INK, width=2.0):
    """A minimal standing figure: head + body line. lean tilts the body."""
    head_d = int(height * 0.20)
    head = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, int(cx - head_d / 2), top, head_d, head_d)
    head.fill.solid()
    head.fill.fore_color.rgb = color
    head.line.fill.background()
    head.shadow.inherit = False
    body_top = top + head_d
    body_h = height - head_d
    offset = int(body_h * lean)
    line(slide, cx, body_top, cx + offset, body_top + body_h,
         color=color, width=width)
    return head


def finish(prs, path, expected=None):
    _cleanup(prs)
    prs.save(path)
    n = len(prs.slides)
    if expected is not None and n != expected:
        raise SystemExit(f"FAIL {path}: {n} slides, expected {expected}")
    return n


def _cleanup(prs):
    for slide in prs.slides:
        _strip_empty_placeholders(slide)
