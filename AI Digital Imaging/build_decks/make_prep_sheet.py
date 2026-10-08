#!/usr/bin/env python3
"""Write PREP_AND_IMAGE_LIST.md from the built lectures.

Pulls the prep notes from each title slide and every image slot (slide number and
what goes in it), so the list can never drift from the files. Run from the repo root
after building the lectures.
"""
import glob
import os
from pptx import Presentation

ROOT = "AI Digital Imaging"
OUT = os.path.join(ROOT, "PREP_AND_IMAGE_LIST.md")

groups = [
    ("AI Digital Imaging lectures", sorted(glob.glob(os.path.join(ROOT, "lectures", "*.pptx")))),
    ("Bridge lectures for Digital Imaging and Painting",
     sorted(glob.glob(os.path.join(ROOT, "bridge_for_Digital_Imaging_and_Painting", "*.pptx")))),
]

lines = [
    "# Prep and image list",
    "",
    "Generated from the lecture files by `build_decks/make_prep_sheet.py`, so it always",
    "matches them. For each lecture: what to make or check before class (from the title",
    "slide's speaker notes), then every image slot, by slide number, with what goes in it.",
    "",
    "Image slots are dashed boxes on the slide. Drag the picture onto the slide, size it",
    "over the box, then delete the box. Anything third-party is meant to be shown in class,",
    "not redistributed; the notes say where a public domain source exists.",
    "",
]

for heading, files in groups:
    lines += [f"## {heading}", ""]
    for f in files:
        prs = Presentation(f)
        title = prs.slides[0].shapes.title.text
        lines += [f"### {title}", "", f"`{os.path.basename(f)}`, {len(prs.slides)} slides", ""]
        note = prs.slides[0].notes_slide.notes_text_frame.text
        for para in note.split("\n"):
            para = para.strip()
            if para.startswith(("Prep", "Check before", "What this is", "Boundary")):
                lines += [para, ""]
        slots = []
        for i, s in enumerate(prs.slides, 1):
            for sh in s.shapes:
                if sh.name == "IMAGE SLOT":
                    slots.append(f"- Slide {i}: {sh.text_frame.text.replace('image: ', '', 1)}")
        if slots:
            lines += ["Image slots:", ""] + slots + [""]

with open(OUT, "w") as fh:
    fh.write("\n".join(lines))
print("wrote", OUT)
