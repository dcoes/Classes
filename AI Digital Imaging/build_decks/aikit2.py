"""Visual slide types for the bridge lecture, built on aikit.

The aim is fewer bullet lists and more built visuals: a timeline, a color-coded
prompt, ladders, a slider, a loop, comparison cards and a big-number slide. All
shapes are native PowerPoint shapes, so they stay editable. Colors come from the
theme (Deck 07): ink 212121, gold D9B247, orange CC702D, red B53A31, olive 7B8865.
"""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from aikit import (L_TITLE_ONLY, notes, visual, _caption, _add_runs, KIND,
                   INK, MUTE, PAPER)

GOLD = RGBColor(0xD9, 0xB2, 0x47)
ORANGE = RGBColor(0xCC, 0x70, 0x2D)
RED = RGBColor(0xB5, 0x3A, 0x31)
OLIVE = RGBColor(0x7B, 0x88, 0x65)
BROWN = RGBColor(0x81, 0x5F, 0x56)
SAND = RGBColor(0xAE, 0x9E, 0x7C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALETTE = [GOLD, ORANGE, RED, OLIVE, BROWN, SAND]

LEFT, RIGHT = 1.42, 11.92
WIDTH = RIGHT - LEFT


def _slide(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE_ONLY])
    s.shapes.title.text = title
    return s


def _text(slide, x, y, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False, space=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _add_runs(p, line)
        p.alignment = align
        p.space_after = Pt(space)
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.italic = italic
            if bold:
                r.font.bold = True
    return tb


def _rect(slide, x, y, w, h, fill=PAPER, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.25):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    return sh


def _line(slide, x1, y1, x2, y2, color=INK, width=2.0, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        from lxml import etree
        tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


# ------------------------------------------------------------------ timeline
def s_timeline(prs, title, events, note, highlight=None):
    """events: list of (year_label, head, sub). highlight: index drawn in red."""
    s = _slide(prs, title)
    y = 4.05
    _line(s, LEFT, y, RIGHT, y, color=SAND, width=3)
    n = len(events)
    step = WIDTH / n
    for i, (year, head, sub) in enumerate(events):
        cx = LEFT + step * (i + 0.5)
        col = RED if highlight == i else GOLD
        _rect(s, cx - 0.13, y - 0.13, 0.26, 0.26, fill=col, shape=MSO_SHAPE.OVAL)
        _text(s, cx - step / 2, y - 1.3, step, 0.5, year, size=28, bold=True,
              color=RED if highlight == i else INK, align=PP_ALIGN.CENTER)
        _text(s, cx - step / 2 + 0.02, y - 0.72, step - 0.04, 0.5, head, size=14, bold=True,
              align=PP_ALIGN.CENTER)
        _text(s, cx - step / 2 + 0.04, y + 0.32, step - 0.08, 1.8, sub, size=13, color=MUTE,
              align=PP_ALIGN.CENTER)
    return notes(s, note)


# ------------------------------------------------------------------ color-coded prompt
def s_prompt(prs, title, parts, note, caption=None):
    """parts: list of (text, label). Each part gets its own color, legend below."""
    s = _slide(prs, title)
    box = _rect(s, LEFT, 2.7, WIDTH, 1.55, fill=RGBColor(0x2A, 0x2A, 0x2A))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.35)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for i, (text, _) in enumerate(parts):
        r = p.add_run()
        r.text = text + ("" if i == len(parts) - 1 else " ")
        r.font.size = Pt(24)
        r.font.color.rgb = PALETTE[i % len(PALETTE)]
    n = len(parts)
    gap = 0.18
    w = (WIDTH - gap * (n - 1)) / n
    for i, (_, label) in enumerate(parts):
        x = LEFT + i * (w + gap)
        _rect(s, x, 4.5, w, 0.1, fill=PALETTE[i % len(PALETTE)])
        _text(s, x, 4.65, w, 0.5, label, size=15, bold=True, align=PP_ALIGN.CENTER)
    if caption:
        _caption(s, caption, 5.55)
    return notes(s, note)


# ------------------------------------------------------------------ ladder
def s_ladder(prs, title, rungs, note, top_label="more control", bottom_label="less control"):
    """rungs: list of (head, sub), lowest first. Rising columns, text above each."""
    s = _slide(prs, title)
    n = len(rungs)
    gap = 0.18
    col_w = (WIDTH - 0.9 - gap * (n - 1)) / n
    base_y = 6.35
    min_h, max_h = 0.3, 1.45
    for i, (head, sub) in enumerate(rungs):
        x = LEFT + 0.9 + i * (col_w + gap)
        h = min_h + (max_h - min_h) * i / max(n - 1, 1)
        y = base_y - h
        shade = [SAND, GOLD, ORANGE, RED, BROWN, OLIVE][i % 6]
        _rect(s, x, y, col_w, h, fill=shade)
        _text(s, x, y - 1.55, col_w, 0.6, head, size=17, bold=True, anchor=MSO_ANCHOR.BOTTOM)
        _text(s, x, y - 0.92, col_w, 0.88, sub, size=13, color=MUTE)
    _text(s, LEFT - 0.05, 2.75, 0.9, 0.6, top_label, size=12, color=MUTE, italic=True)
    _text(s, LEFT - 0.05, 5.9, 0.9, 0.5, bottom_label, size=12, color=MUTE, italic=True)
    _line(s, LEFT + 0.35, 5.8, LEFT + 0.35, 3.45, color=MUTE, width=1.5, arrow=True)
    return notes(s, note)


# ------------------------------------------------------------------ slider
def s_slider(prs, title, left_label, right_label, stops, note, caption=None):
    """stops: list of (position 0..1, head, sub)."""
    s = _slide(prs, title)
    y = 3.75
    bar = _rect(s, LEFT, y - 0.09, WIDTH, 0.18, fill=SAND)
    _text(s, LEFT, y - 0.85, 5, 0.5, left_label, size=18, bold=True)
    _text(s, RIGHT - 5, y - 0.85, 5, 0.5, right_label, size=18, bold=True, align=PP_ALIGN.RIGHT)
    for pos, head, sub in stops:
        cx = LEFT + pos * WIDTH
        _rect(s, cx - 0.17, y - 0.17, 0.34, 0.34, fill=RED, shape=MSO_SHAPE.OVAL)
        _text(s, cx - 1.55, y + 0.3, 3.1, 0.45, head, size=19, bold=True, align=PP_ALIGN.CENTER)
        _text(s, cx - 1.55, y + 0.78, 3.1, 1.4, sub, size=15, color=MUTE, align=PP_ALIGN.CENTER)
    if caption:
        _caption(s, caption, 5.9)
    return notes(s, note)


# ------------------------------------------------------------------ loop
def s_loop(prs, title, stages, note, human=None, caption=None, ret=True):
    """stages: list of (head, sub) left to right, optional return arrow underneath.
    human: indexes drawn in red (the steps only a person can do)."""
    s = _slide(prs, title)
    human = human or []
    n = len(stages)
    gap = 0.38
    w = (WIDTH - gap * (n - 1)) / n
    y = 2.85
    h = 1.75
    for i, (head, sub) in enumerate(stages):
        x = LEFT + i * (w + gap)
        col = RED if i in human else RGBColor(0x2A, 0x2A, 0x2A)
        _rect(s, x, y, w, h, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        _text(s, x + 0.06, y + 0.12, w - 0.12, 0.85, head, size=17 if n <= 5 else 15, bold=True,
              color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        _text(s, x + 0.06, y + 1.0, w - 0.12, 0.7, sub, size=13, color=WHITE, align=PP_ALIGN.CENTER)
        if i < n - 1:
            _line(s, x + w + 0.03, y + h / 2, x + w + gap - 0.03, y + h / 2, color=INK, width=2, arrow=True)
    if ret:
        ry = y + h + 0.4
        first_cx = LEFT + w / 2
        last_cx = LEFT + (n - 1) * (w + gap) + w / 2
        _line(s, last_cx, y + h, last_cx, ry, color=MUTE, width=1.5)
        _line(s, last_cx, ry, first_cx, ry, color=MUTE, width=1.5)
        _line(s, first_cx, ry, first_cx, y + h + 0.02, color=MUTE, width=1.5, arrow=True)
        _text(s, LEFT, ry + 0.05, WIDTH, 0.4, "not right yet? back around", size=13, color=MUTE,
              italic=True, align=PP_ALIGN.CENTER)
    if caption:
        _caption(s, caption, 5.95)
    return notes(s, note)


# ------------------------------------------------------------------ cards
def s_cards(prs, title, cards, note, caption=None):
    """cards: list of (head, [lines]) shown as 2 to 4 cards with a colored header band."""
    s = _slide(prs, title)
    n = len(cards)
    gap = 0.3
    w = (WIDTH - gap * (n - 1)) / n
    y = 2.75
    h = 3.25 if not caption else 2.95
    for i, (head, lines) in enumerate(cards):
        x = LEFT + i * (w + gap)
        _rect(s, x, y, w, h, fill=PAPER, line=SAND, lw=1)
        band = _rect(s, x, y, w, 0.6, fill=PALETTE[i % len(PALETTE)])
        _text(s, x + 0.15, y + 0.06, w - 0.3, 0.5, head, size=18 if n <= 3 else 16, bold=True,
              color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        body = {1: 22, 2: 21, 3: 18}.get(n, 15)
        avail = h - 1.0
        while body > 12:
            cpl = max(1, int((w - 0.45) * 72 / (0.52 * body)))
            nlines = sum(max(1, -(-len(t) // cpl)) for t in lines)
            need = nlines * body * 1.25 / 72 + len(lines) * 8 / 72
            if need <= avail:
                break
            body -= 1
        _text(s, x + 0.2, y + 0.8, w - 0.4, h - 0.9, "\n".join(lines), size=body, space=8)
    if caption:
        _caption(s, caption, 5.85)
    return notes(s, note)


# ------------------------------------------------------------------ rows
def s_rows(prs, title, rows, note):
    """rows: list of (label, text). A colored label chip, then the text, one row each."""
    s = _slide(prs, title)
    n = len(rows)
    row_h = min(0.72, 3.7 / n)
    for i, (label, text) in enumerate(rows):
        y = 2.75 + i * row_h
        _rect(s, LEFT, y + 0.06, 2.6, row_h - 0.14, fill=PALETTE[i % len(PALETTE)])
        _text(s, LEFT + 0.15, y + 0.06, 2.4, row_h - 0.14, label, size=19, bold=True, color=WHITE,
              anchor=MSO_ANCHOR.MIDDLE)
        _text(s, LEFT + 2.85, y + 0.06, WIDTH - 2.9, row_h - 0.14, text, size=19,
              anchor=MSO_ANCHOR.MIDDLE)
    return notes(s, note)


# ------------------------------------------------------------------ big number
def s_bignum(prs, title, number, line1, line2, note, source=None):
    s = _slide(prs, title)
    _text(s, LEFT, 2.65, 4.2, 2.2, number, size=96, bold=True, color=RED, anchor=MSO_ANCHOR.MIDDLE)
    _text(s, LEFT + 4.4, 2.95, WIDTH - 4.4, 1.0, line1, size=22, bold=True)
    _text(s, LEFT + 4.4, 3.95, WIDTH - 4.4, 1.6, line2, size=16, color=MUTE)
    if source:
        _text(s, LEFT, 5.95, WIDTH, 0.4, source, size=11, color=MUTE, italic=True)
    return notes(s, note)


# ------------------------------------------------------------------ statement
def s_statement(prs, title, line, sub, note):
    """One idea, large, no quote marks. For a sentence the lecture hinges on."""
    s = _slide(prs, title)
    _rect(s, LEFT, 2.9, 0.12, 2.9, fill=GOLD)
    _text(s, LEFT + 0.4, 2.8, WIDTH - 0.5, 1.5, line, size=34, anchor=MSO_ANCHOR.MIDDLE)
    _text(s, LEFT + 0.4, 4.4, WIDTH - 0.5, 1.6, sub, size=19, color=MUTE)
    return notes(s, note)


# ------------------------------------------------------------------ LoRA diagram
def s_lora(prs, title, note):
    s = _slide(prs, title)
    # base model block
    _rect(s, LEFT, 2.85, 4.3, 2.6, fill=RGBColor(0x2A, 0x2A, 0x2A), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _text(s, LEFT + 0.2, 3.0, 3.9, 0.5, "The base model", size=20, bold=True, color=WHITE)
    _text(s, LEFT + 0.2, 3.55, 3.9, 1.8,
          "Several gigabytes.\nTrained on millions of images.\nKnows a little about everything.",
          size=17, color=WHITE)
    # LoRA clip-on
    _rect(s, LEFT + 4.75, 3.35, 2.3, 1.6, fill=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _text(s, LEFT + 4.9, 3.45, 2.0, 0.5, "A LoRA", size=20, bold=True, color=WHITE)
    _text(s, LEFT + 4.9, 3.95, 2.0, 1.0, "A small add-on file.\nTrained on 20 to 50 images.", size=14, color=WHITE)
    _line(s, LEFT + 4.32, 4.15, LEFT + 4.72, 4.15, color=INK, width=2.5)
    # inputs and result
    _text(s, LEFT + 7.45, 2.85, 3.0, 0.4, "You add", size=13, color=MUTE, italic=True)
    _text(s, LEFT + 7.45, 3.2, 3.05, 1.0, '"mara_character"\na trigger word in the prompt', size=15, bold=True)
    _text(s, LEFT + 7.45, 4.15, 3.0, 0.4, "and a strength", size=13, color=MUTE, italic=True)
    _rect(s, LEFT + 7.45, 4.6, 2.8, 0.12, fill=SAND)
    _rect(s, LEFT + 7.45, 4.6, 2.0, 0.12, fill=RED)
    _text(s, LEFT + 7.45, 4.75, 3.0, 0.5, "0.7 of full strength", size=12, color=MUTE)
    _text(s, LEFT, 5.75, WIDTH, 0.5,
          "Like a clip-on lens: the camera stays the same, one thing about every picture changes.",
          size=15, align=PP_ALIGN.CENTER)
    return notes(s, note)


KIND.update({
    "timeline": s_timeline, "prompt": s_prompt, "ladder": s_ladder, "slider": s_slider,
    "loop": s_loop, "cards": s_cards, "bignum": s_bignum, "statement": s_statement,
    "lora": s_lora, "rows": s_rows,
})
